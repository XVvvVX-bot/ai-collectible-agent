from __future__ import annotations

import csv
import sqlite3
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from ai_agent.storage.sqlite_raw_store import SqliteRawStore
from ai_agent.user_domain import UserDomainService, canonical_item_dedupe_key

HEADER_ALIASES = {
    "item_type": "item_type",
    "type": "item_type",
    "holding_or_watch": "item_type",
    "list_type": "item_type",
    "category": "category",
    "interest_category": "category",
    "series": "series",
    "interest_series": "series",
    "item_name": "item_name",
    "name": "item_name",
    "title": "item_name",
    "year": "year",
    "grade_condition": "grade_condition",
    "grade": "grade_condition",
    "grade_target": "grade_condition",
    "condition": "grade_condition",
    "quantity": "quantity",
    "qty": "quantity",
    "cost_basis_total": "cost_basis_total",
    "cost": "cost_basis_total",
    "total_cost": "cost_basis_total",
    "priority": "priority",
    "max_buy_price": "max_buy_price",
    "budget": "max_buy_price",
    "notes": "notes",
    "is_active": "is_active",
    "active": "is_active",
}

@dataclass(frozen=True)
class ImportErrorDetail:
    row_number: int
    code: str
    message: str


@dataclass(frozen=True)
class UserItemsImportResult:
    mode: str
    file_path: str
    total_rows: int
    inserted: int
    updated: int
    skipped: int
    errors: list[ImportErrorDetail]

    def to_dict(self) -> dict[str, Any]:
        return {
            "mode": self.mode,
            "file_path": self.file_path,
            "total_rows": self.total_rows,
            "inserted": self.inserted,
            "updated": self.updated,
            "skipped": self.skipped,
            "errors": [asdict(err) for err in self.errors],
        }


def import_user_items_file(
    db_path: str,
    user_id: str,
    file_path: str,
    *,
    commit: bool = False,
) -> UserItemsImportResult:
    source_path = Path(file_path)
    rows = _load_rows(source_path)
    service = UserDomainService(db_path)
    SqliteRawStore(db_path).ensure_schema()
    if commit:
        service.upsert_user(user_id)

    total_rows = len(rows)
    inserted = 0
    updated = 0
    skipped = 0
    errors: list[ImportErrorDetail] = []
    seen_batch_keys: set[tuple[str, str]] = set()

    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        for row_number, row in rows:
            if _row_is_blank(row):
                skipped += 1
                continue
            try:
                parsed = _parse_row(row)
            except ValueError as exc:
                errors.append(
                    ImportErrorDetail(
                        row_number=row_number,
                        code="invalid_row",
                        message=str(exc),
                    )
                )
                continue

            item_type = parsed["item_type"]
            dedupe_key = canonical_item_dedupe_key(
                category=parsed["category"],
                series=parsed["series"],
                item_name=parsed["item_name"],
                year=parsed["year"],
                grade_condition=parsed["grade_condition"],
            )
            batch_key = (item_type, dedupe_key)
            if batch_key in seen_batch_keys:
                errors.append(
                    ImportErrorDetail(
                        row_number=row_number,
                        code="duplicate_in_file",
                        message="Row duplicates a prior row in this file by natural key.",
                    )
                )
                continue
            seen_batch_keys.add(batch_key)

            exists_before = _item_exists(conn, user_id=user_id, item_type=item_type, dedupe_key=dedupe_key)
            if commit:
                service.upsert_user_item(
                    user_id,
                    item_type,
                    category=parsed["category"],
                    series=parsed["series"],
                    item_name=parsed["item_name"],
                    year=parsed["year"],
                    grade_condition=parsed["grade_condition"],
                    quantity=parsed["quantity"],
                    cost_basis_total=parsed["cost_basis_total"],
                    priority=parsed["priority"],
                    max_buy_price=parsed["max_buy_price"],
                    notes=parsed["notes"],
                    is_active=parsed["is_active"],
                )
            if exists_before:
                updated += 1
            else:
                inserted += 1

    return UserItemsImportResult(
        mode="commit" if commit else "preview",
        file_path=str(source_path),
        total_rows=total_rows,
        inserted=inserted,
        updated=updated,
        skipped=skipped,
        errors=errors,
    )


def _item_exists(
    conn: sqlite3.Connection,
    *,
    user_id: str,
    item_type: str,
    dedupe_key: str,
) -> bool:
    row = conn.execute(
        """
        SELECT 1
        FROM user_items
        WHERE user_id = ? AND item_type = ? AND dedupe_key = ?
        """,
        (user_id, item_type, dedupe_key),
    ).fetchone()
    return row is not None


def _load_rows(file_path: Path) -> list[tuple[int, dict[str, Any]]]:
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        return _load_csv_rows(file_path)
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx_rows(file_path)
    raise ValueError(f"Unsupported file extension: {file_path.suffix}. Use .csv or .xlsx.")


def _load_csv_rows(file_path: Path) -> list[tuple[int, dict[str, Any]]]:
    if not file_path.exists():
        raise ValueError(f"Input file not found: {file_path}")
    rows: list[tuple[int, dict[str, Any]]] = []
    with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError("CSV file is missing a header row.")
        for idx, row in enumerate(reader, start=2):
            rows.append((idx, dict(row)))
    return rows


def _load_xlsx_rows(file_path: Path) -> list[tuple[int, dict[str, Any]]]:
    if not file_path.exists():
        raise ValueError(f"Input file not found: {file_path}")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel import requires openpyxl. Install it or use CSV.") from exc

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    header_values: list[str] = []
    rows: list[tuple[int, dict[str, Any]]] = []
    for row_idx, cells in enumerate(ws.iter_rows(values_only=True), start=1):
        values = ["" if value is None else str(value) for value in cells]
        if row_idx == 1:
            header_values = values
            continue
        row_map: dict[str, Any] = {}
        for col_idx, raw_header in enumerate(header_values):
            if not raw_header:
                continue
            value = values[col_idx] if col_idx < len(values) else ""
            row_map[raw_header] = value
        rows.append((row_idx, row_map))
    if not header_values:
        raise ValueError("Excel file is missing a header row.")
    return rows


def _parse_row(row: dict[str, Any]) -> dict[str, Any]:
    canonical = _normalize_headers(row)
    item_type = _parse_item_type(canonical.get("item_type"))
    item_name = _required_text(canonical.get("item_name"), "item_name")
    year = _optional_int(canonical.get("year"), "year")
    quantity = _optional_float(canonical.get("quantity"), "quantity")
    cost_basis_total = _optional_float(canonical.get("cost_basis_total"), "cost_basis_total")
    max_buy_price = _optional_float(canonical.get("max_buy_price"), "max_buy_price")
    is_active = _optional_bool(canonical.get("is_active"), "is_active", default=True)
    priority = _optional_text(canonical.get("priority"))
    if priority is not None:
        priority = priority.lower()
        if priority not in {"high", "normal", "low"}:
            raise ValueError(f"priority must be one of high/normal/low; got '{priority}'.")

    if item_type == "watch" and quantity is not None:
        raise ValueError("quantity is not allowed for watch items.")
    if item_type == "holding" and priority is not None:
        raise ValueError("priority is not allowed for holding items.")

    return {
        "item_type": item_type,
        "category": _optional_text(canonical.get("category")),
        "series": _optional_text(canonical.get("series")),
        "item_name": item_name,
        "year": year,
        "grade_condition": _optional_text(canonical.get("grade_condition")),
        "quantity": quantity,
        "cost_basis_total": cost_basis_total,
        "priority": priority,
        "max_buy_price": max_buy_price,
        "notes": _optional_text(canonical.get("notes")),
        "is_active": is_active,
    }


def _normalize_headers(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for raw_key, value in row.items():
        key = _normalize_header_key(raw_key)
        canonical = HEADER_ALIASES.get(key)
        if canonical is None:
            continue
        normalized[canonical] = value
    return normalized


def _normalize_header_key(value: Any) -> str:
    text = _optional_text(value) or ""
    text = text.replace("-", "_").replace(" ", "_")
    return text.lower()


def _parse_item_type(value: Any) -> str:
    text = (_optional_text(value) or "").lower()
    if text in {"holding", "hold", "h"}:
        return "holding"
    if text in {"watch", "watchlist", "w"}:
        return "watch"
    raise ValueError(f"item_type is required and must be holding/watch; got '{text}'.")


def _required_text(value: Any, field_name: str) -> str:
    text = _optional_text(value)
    if not text:
        raise ValueError(f"{field_name} is required.")
    return text


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _optional_int(value: Any, field_name: str) -> int | None:
    text = _optional_text(value)
    if text is None:
        return None
    try:
        return int(text)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be an integer; got '{text}'.") from exc


def _optional_float(value: Any, field_name: str) -> float | None:
    text = _optional_text(value)
    if text is None:
        return None
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be a number; got '{text}'.") from exc


def _optional_bool(value: Any, field_name: str, *, default: bool) -> bool:
    text = _optional_text(value)
    if text is None:
        return default
    lowered = text.lower()
    if lowered in {"1", "true", "yes", "y"}:
        return True
    if lowered in {"0", "false", "no", "n"}:
        return False
    raise ValueError(f"{field_name} must be a boolean value; got '{text}'.")


def _row_is_blank(row: dict[str, Any]) -> bool:
    for value in row.values():
        if _optional_text(value):
            return False
    return True
